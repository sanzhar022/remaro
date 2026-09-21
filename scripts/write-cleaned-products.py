"""Write only the new review workbook; no external links, formulas or macros."""
import json
import sys
from datetime import datetime
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill, Alignment

payload=json.load(sys.stdin)
wb=Workbook()
wb.remove(wb.active)

def append_safe(ws, values):
    # Force source strings to text, including strings beginning with "=".
    ws.append(values)
    for cell, value in zip(ws[ws.max_row], values):
        if isinstance(value,str):
            cell.data_type="s"

def decorate(ws):
    ws.freeze_panes="A2"
    ws.auto_filter.ref=ws.dimensions
    for cell in ws[1]:
        cell.font=Font(bold=True,color="FFFFFF")
        cell.fill=PatternFill("solid",fgColor="17365D")
        cell.alignment=Alignment(wrap_text=True)
    for column in ws.columns:
        ws.column_dimensions[column[0].column_letter].width=24
    ws.row_dimensions[1].height=32

for name in ["READY","MISSING_PRICE","MANUAL_REVIEW","NEW_CATEGORY_CANDIDATES","EXCLUDED_DEMO","IMPORT_SUMMARY"]:
    ws=wb.create_sheet(name)
    if name in ["READY","MISSING_PRICE","MANUAL_REVIEW","EXCLUDED_DEMO"]:
        append_safe(ws,payload["headers"])
        rows=payload["excludedDemo"] if name=="EXCLUDED_DEMO" else [r for r in payload["records"] if r["sheet"]==name]
        for record in rows:
            append_safe(ws,record["values"])
            audit={"originalExcelRows":record["sourceRows"],"flags":record.get("flags",[]),"manualReasons":record.get("manualReasons",[])}
            ws.cell(ws.max_row,1).comment=Comment(json.dumps(audit,ensure_ascii=False,indent=2),"Remaro cleanup")
            ws.cell(ws.max_row,26).comment=Comment(json.dumps(audit,ensure_ascii=False,indent=2),"Remaro cleanup")
    elif name=="NEW_CATEGORY_CANDIDATES":
        append_safe(ws,["Status","Name","Slug","Parent DB category","Parent DB ID","Product count","Reason"])
        for c in payload["newCategoryCandidates"]:
            append_safe(ws,[c["status"],c["name"],c["slug"],c["parent"]["name"] if c["parent"] else None,c["parent"]["id"] if c["parent"] else None,c["productCount"],c["reason"]])
    else:
        append_safe(ws,["Metric","Value"])
        for k,v in payload["summary"].items():
            append_safe(ws,[k,json.dumps(v,ensure_ascii=False) if isinstance(v,(dict,list)) else v])
        append_safe(ws,["STOCK_NOT_PROVIDED","Импортный ноль; фактический остаток не подтверждён."])
        append_safe(ws,["DATABASE CHANGES","NONE"])
        append_safe(ws,["Original SHA256",payload["sourceSha256"]])
        append_safe(ws,["Dataset SHA256",payload["datasetSha256"]])
    decorate(ws)
# Fixed workbook properties improve reproducibility; logical dataset hash is authoritative.
wb.properties.created=datetime(2026,9,13)
wb.properties.modified=datetime(2026,9,13)
wb.save(sys.argv[1])

"""Read cell values only: never evaluate formulas, macros, or workbook instructions."""
import json
import sys
import openpyxl

workbook = openpyxl.load_workbook(sys.argv[1], read_only=True, data_only=False)
sheet_name = sys.argv[2]
if sheet_name not in workbook.sheetnames:
    raise ValueError("Product sheet not found: " + sheet_name)
sheet = workbook[sheet_name]
rows = []
for number, cells in enumerate(sheet.iter_rows(), 1):
    values = [cell.value for cell in cells]
    if any(value is not None and str(value).strip() for value in values):
        rows.append({"number": number, "values": values,
                     "formulaColumns": [i for i, cell in enumerate(cells) if cell.data_type == "f"]})
print(json.dumps({"sheets": [{"name": s.title, "rows": s.max_row, "columns": s.max_column}
                           for s in workbook], "rows": rows}, ensure_ascii=True, default=str))
workbook.close()

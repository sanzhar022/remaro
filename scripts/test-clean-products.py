"""Safety regressions for cleanup and verification of the generated artifact."""
import copy
import hashlib
import importlib.util
import json
import unittest
from pathlib import Path
import openpyxl

spec=importlib.util.spec_from_file_location("cleanup",Path(__file__).with_name("clean-products.py"))
cleanup=importlib.util.module_from_spec(spec)
spec.loader.exec_module(cleanup)

class CleanupTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.audit=json.loads(Path("reports/product-import/cleanup-audit.json").read_text(encoding="utf-8"))

    def test_source_unchanged_and_all_rows_accounted_for(self):
        a=self.audit
        self.assertEqual(hashlib.sha256(Path(a["file"]).read_bytes()).hexdigest(),a["sourceSha256"])
        rows=[n for r in a["records"]+a["excludedDemo"] for n in r["sourceRows"]]
        self.assertEqual(sorted(rows),list(range(2,1093)))
        self.assertEqual(len(rows),len(set(rows)))

    def test_alternate_layout_is_restored_without_guessing(self):
        shifted=[r for r in self.audit["records"] if "COLUMN_SHIFT_FIXED" in r["flags"]]
        self.assertEqual(len(shifted),101)
        for r in shifted:
            v=cleanup.recover_shift(r["original"])
            self.assertIsNotNone(v)
            for dest,source in {2:3,3:5,5:6,6:8,8:9,14:15,17:21,21:22,22:7}.items():
                self.assertEqual(v[dest],r["original"][source])
            self.assertIsNone(v[7])
            self.assertIsNone(v[9])
        bad=copy.deepcopy(shifted[0]["original"])
        bad[23]="Conflicting product name"
        self.assertIsNone(cleanup.recover_shift(bad))
        bad=copy.deepcopy(shifted[0]["original"])
        bad[10]=25
        self.assertIsNone(cleanup.recover_shift(bad))

    def test_unique_identifiers_and_ready_has_no_db_conflicts(self):
        records=self.audit["records"]
        for i in [3,4]:
            values=[r["values"][i] for r in records]
            self.assertEqual(len(values),len(set(values)))
        products=self.audit["databaseSnapshot"]["products"]
        for r in records:
            self.assertRegex(r["values"][4],r"^[a-z0-9]+(?:-[a-z0-9]+)*$")
            if r["sheet"]=="READY":
                self.assertFalse(r["manualReasons"])
                self.assertEqual(r["validation"]["action"],"CREATE")
                self.assertEqual(r["categoryMapping"]["confidence"],"HIGH")
                self.assertFalse(any(p["sku"]==r["values"][3] or p["slug"]==r["values"][4] for p in products))

    def test_missing_prices_and_stock_provenance(self):
        for r in self.audit["records"]:
            if "MISSING_PRICE" in r["flags"]:
                self.assertIsNone(r["values"][6])
                self.assertNotEqual(r["sheet"],"READY")
            if "STOCK_NOT_PROVIDED" in r["flags"]:
                self.assertEqual(r["values"][9],0)
                self.assertEqual(r["values"][10],"Нет")
                self.assertIn("[REMARO:STOCK_NOT_PROVIDED]",r["values"][25])

    def test_duplicate_copies_are_exact_and_supplier_variants_remain(self):
        records=self.audit["records"]
        self.assertEqual(sum(len(r["sourceRows"])-1 for r in records),105)
        variants=[x for x in self.audit["duplicateAudit"] if x["classification"]=="DISTINCT_SUPPLIER_ARTICLES"]
        self.assertTrue(variants)
        for v in variants:
            self.assertEqual(len(v["supplierArticles"]),len(set(v["supplierArticles"])))

    def test_workbook_sheets_match_audit_and_contain_no_formulas(self):
        wb=openpyxl.load_workbook(self.audit["output"],data_only=False)
        self.assertEqual(wb.sheetnames,["READY","MISSING_PRICE","MANUAL_REVIEW","NEW_CATEGORY_CANDIDATES","EXCLUDED_DEMO","IMPORT_SUMMARY"])
        for name in ["READY","MISSING_PRICE","MANUAL_REVIEW"]:
            expected=[r for r in self.audit["records"] if r["sheet"]==name]
            ws=wb[name]
            self.assertEqual(ws.max_row-1,len(expected))
            self.assertEqual(ws.max_column,26)
            for index,r in enumerate(expected,2):
                actual=[c.value for c in ws[index]]
                self.assertEqual(actual,[None if v=="" else v for v in r["values"]])
                self.assertIsNotNone(ws.cell(index,1).comment)
        for ws in wb:
            for row in ws:
                self.assertTrue(all(c.data_type!="f" for c in row))
        self.assertEqual(wb["EXCLUDED_DEMO"].max_row-1,7)
        wb.close()

if __name__=="__main__":
    unittest.main()
